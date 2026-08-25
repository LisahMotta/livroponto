from openpyxl import load_workbook

from livroponto.models import DiaNaoLetivo, Escola, Licenca, LivroPontoConfig, Pessoa, TipoServidor
from livroponto.readers.template_reader import criar_modelo, ler_modelo, salvar_modelo


def test_criar_modelo_gera_abas_esperadas(tmp_path):
    caminho = tmp_path / "modelo.xlsx"
    criar_modelo(caminho)
    wb = load_workbook(caminho)
    assert "Escola" in wb.sheetnames
    assert "Pessoas" in wb.sheetnames


def test_ler_modelo_roundtrip(tmp_path):
    caminho = tmp_path / "modelo.xlsx"
    criar_modelo(caminho)

    wb = load_workbook(caminho)
    aba_escola = wb["Escola"]
    for row in aba_escola.iter_rows(min_row=2):
        if row[0].value == "nome":
            row[1].value = "EE Exemplo Fictício"
        if row[0].value == "municipio":
            row[1].value = "Cidade Exemplo"
        if row[0].value == "mes":
            row[1].value = 4
        if row[0].value == "ano":
            row[1].value = 2026
    wb.save(caminho)

    config = ler_modelo(caminho)
    assert config.escola.nome == "EE Exemplo Fictício"
    assert config.mes == 4
    assert config.ano == 2026
    assert len(config.pessoas) == 3
    tipos = {p.tipo for p in config.pessoas}
    assert tipos == {TipoServidor.ADMINISTRATIVO, TipoServidor.DOCENTE, TipoServidor.GESTAO}
    assert config.nome_diretor() == "Beltrano de Souza"


def test_salvar_modelo_e_ler_modelo_preservam_excecoes_e_jornada_codigo(tmp_path):
    config = LivroPontoConfig(
        escola=Escola(nome="EE Exemplo Fictício", municipio="Cidade Exemplo"),
        mes=7,
        ano=2026,
        uf="SP",
        cidade_assinatura="Cidade Exemplo",
        pessoas=[
            Pessoa(
                nome="Docente Fictício",
                tipo=TipoServidor.DOCENTE,
                rg="1.111.111-1",
                cargo="PEB II",
                jornada_codigo="B",
                disciplinas="HISTÓRIA",
            ),
        ],
        dias_excecao=[
            DiaNaoLetivo(mes=7, dia=8, tipo="RECESSO", descricao="Recesso escolar"),
        ],
    )

    caminho = tmp_path / "modelo.xlsx"
    salvar_modelo(config, caminho)
    recarregado = ler_modelo(caminho)

    assert recarregado.uf == "SP"
    assert len(recarregado.pessoas) == 1
    assert recarregado.pessoas[0].jornada_codigo == "B"
    assert len(recarregado.dias_excecao) == 1
    assert recarregado.dias_excecao[0].tipo == "RECESSO"
    assert recarregado.dias_excecao[0].dia == 8


def test_salvar_modelo_e_ler_modelo_preservam_regime_plantao_e_horario_estudante(tmp_path):
    config = LivroPontoConfig(
        escola=Escola(nome="EE Exemplo Fictício", municipio="Cidade Exemplo"),
        mes=7,
        ano=2026,
        pessoas=[
            Pessoa(
                nome="Servidor Fictício",
                tipo=TipoServidor.ADMINISTRATIVO,
                rg="1.111.111-1",
                cargo="Agente de Organização Escolar",
                regime_plantao="Sim",
                horario_estudante="Não",
            ),
        ],
    )

    caminho = tmp_path / "modelo.xlsx"
    salvar_modelo(config, caminho)
    recarregado = ler_modelo(caminho)

    assert recarregado.pessoas[0].regime_plantao == "Sim"
    assert recarregado.pessoas[0].horario_estudante == "Não"


def test_salvar_modelo_e_ler_modelo_preservam_trio_gestor(tmp_path):
    config = LivroPontoConfig(
        escola=Escola(nome="EE Exemplo Fictício", municipio="Cidade Exemplo"),
        mes=4,
        ano=2026,
        diretor_nome="Diretora Fictícia",
        pessoas=[
            Pessoa(
                nome="Diretora Fictícia",
                tipo=TipoServidor.GESTAO,
                cargo="Diretor(a) de Escola",
                rg="1.234.567-8",
                jornada_semanal=40,
                entrada="07:00",
                saida="16:00",
            ),
            Pessoa(nome="Vice Fictício", tipo=TipoServidor.GESTAO, cargo="Vice-Diretor(a) de Escola"),
        ],
    )

    caminho = tmp_path / "modelo.xlsx"
    salvar_modelo(config, caminho)
    recarregado = ler_modelo(caminho)

    gestao = recarregado.pessoas_por_tipo(TipoServidor.GESTAO)
    assert len(gestao) == 2
    assert gestao[0].nome == "Diretora Fictícia"
    assert gestao[0].rg == "1.234.567-8"
    assert gestao[0].jornada_semanal == 40
    assert gestao[0].entrada == "07:00"
    assert recarregado.nome_diretor() == "Diretora Fictícia"


def test_salvar_modelo_e_ler_modelo_preservam_periodo_de_ferias(tmp_path):
    config = LivroPontoConfig(
        escola=Escola(nome="EE Exemplo Fictício", municipio="Cidade Exemplo"),
        mes=4,
        ano=2026,
        pessoas=[
            Pessoa(
                nome="Servidor Fictício",
                tipo=TipoServidor.ADMINISTRATIVO,
                rg="1.111.111-1",
                ferias_inicio="03/04/2026",
                ferias_fim="02/05/2026",
            ),
        ],
    )

    caminho = tmp_path / "modelo.xlsx"
    salvar_modelo(config, caminho)
    recarregado = ler_modelo(caminho)

    assert recarregado.pessoas[0].ferias_inicio == "03/04/2026"
    assert recarregado.pessoas[0].ferias_fim == "02/05/2026"
    assert recarregado.pessoas[0].periodo_ferias == "03/04/2026 a 02/05/2026"


def test_salvar_modelo_e_ler_modelo_preservam_licencas(tmp_path):
    config = LivroPontoConfig(
        escola=Escola(nome="EE Exemplo Fictício", municipio="Cidade Exemplo"),
        mes=4,
        ano=2026,
        pessoas=[
            Pessoa(
                nome="Servidor Fictício",
                tipo=TipoServidor.ADMINISTRATIVO,
                rg="1.111.111-1",
                licencas=[
                    Licenca(tipo="SAUDE", inicio="10/04/2026", fim="20/04/2026"),
                    Licenca(tipo="PREMIO", inicio="01/06/2026", fim="30/06/2026"),
                ],
            ),
            Pessoa(nome="Sem Licença", tipo=TipoServidor.ADMINISTRATIVO, rg="2.222.222-2"),
        ],
    )

    caminho = tmp_path / "modelo.xlsx"
    salvar_modelo(config, caminho)
    recarregado = ler_modelo(caminho)

    com_licenca = next(p for p in recarregado.pessoas if p.nome == "Servidor Fictício")
    sem_licenca = next(p for p in recarregado.pessoas if p.nome == "Sem Licença")
    assert len(com_licenca.licencas) == 2
    assert {(lic.tipo, lic.periodo) for lic in com_licenca.licencas} == {
        ("SAUDE", "10/04/2026 a 20/04/2026"),
        ("PREMIO", "01/06/2026 a 30/06/2026"),
    }
    assert sem_licenca.licencas == []


def test_salvar_modelo_e_ler_modelo_preservam_rotulo_de_assinatura(tmp_path):
    config = LivroPontoConfig(
        escola=Escola(nome="URE Exemplo", municipio="Cidade Exemplo"),
        mes=4,
        ano=2026,
        rotulo_assinatura="Dirigente Regional de Ensino",
    )

    caminho = tmp_path / "modelo.xlsx"
    salvar_modelo(config, caminho)
    recarregado = ler_modelo(caminho)

    assert recarregado.rotulo_assinatura == "Dirigente Regional de Ensino"
    assert recarregado.rotulo_assinatura_efetivo() == "Dirigente Regional de Ensino"
