"""Leitor/gerador do modelo simplificado (.xlsx) — alternativa para quem não
tem a planilha legada SEDUC-SP: uma aba "Escola" (campo/valor), uma aba
"Pessoas" (uma linha por servidor — administrativo, docente ou do trio
gestor, todos na mesma aba, diferenciados pela coluna "tipo") e uma aba
"Excecoes" (recesso, ponto facultativo etc.). É o mesmo formato usado para
persistir o que é editado no app web (`livroponto app`) — editar lá e
"Salvar" grava nesse formato."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from ..models import DiaNaoLetivo, Escola, LivroPontoConfig, Pessoa, TipoServidor

COLUNAS_PESSOAS = [
    "tipo",
    "nome",
    "rg",
    "cargo",
    "jornada_semanal",
    "jornada_codigo",
    "ponto",
    "entrada",
    "saida",
    "intervalo_inicio",
    "intervalo_fim",
    "ferias_inicio",
    "ferias_fim",
    "disciplinas",
    "categoria",
    "situacao",
    "observacoes",
]

COLUNAS_EXCECOES = ["mes", "dia", "tipo", "descricao"]

CAMPOS_ESCOLA = [
    ("nome", "Nome da escola"),
    ("diretoria_ensino", "Diretoria de Ensino"),
    ("endereco", "Endereço"),
    ("municipio", "Município"),
    ("telefone1", "Telefone 1"),
    ("telefone2", "Telefone 2"),
    ("email", "E-mail"),
    ("codigo_ua", "Código UA"),
    ("codigo_cie", "Código CIE"),
]

CAMPOS_EXTRAS = ["mes", "ano", "cidade_assinatura", "diretor_nome", "uf"]


def _linha_pessoa(p: Pessoa) -> list:
    return [
        p.tipo.value,
        p.nome,
        p.rg,
        p.cargo,
        p.jornada_semanal if p.jornada_semanal is not None else "",
        p.jornada_codigo,
        "S" if p.ponto else "N",
        p.entrada,
        p.saida,
        p.intervalo_inicio,
        p.intervalo_fim,
        p.ferias_inicio,
        p.ferias_fim,
        p.disciplinas,
        p.categoria,
        p.situacao,
        p.observacoes,
    ]


def criar_modelo(caminho: str | Path) -> None:
    """Gera um arquivo .xlsx em branco (com um exemplo) para o usuário
    preencher, no formato aceito por `ler_modelo`."""
    escola = Escola()
    pessoas = [
        Pessoa(
            nome="Fulano de Tal",
            tipo=TipoServidor.ADMINISTRATIVO,
            rg="12.345.678-9",
            cargo="Agente de Organização Escolar",
            jornada_semanal=40,
            entrada="07:00",
            saida="16:00",
            intervalo_inicio="11:00",
            intervalo_fim="12:00",
        ),
        Pessoa(
            nome="Ciclana da Silva",
            tipo=TipoServidor.DOCENTE,
            rg="98.765.432-1",
            cargo="Professor de Educação Básica II",
            categoria="Titular de Cargo",
            disciplinas="MATEMÁTICA",
            jornada_codigo="B",
        ),
        Pessoa(
            nome="Beltrano de Souza",
            tipo=TipoServidor.GESTAO,
            rg="11.222.333-4",
            cargo="Diretor(a) de Escola",
            jornada_semanal=40,
            entrada="07:00",
            saida="16:00",
        ),
    ]
    config = LivroPontoConfig(escola=escola, mes=4, ano=2026, pessoas=pessoas, diretor_nome="Beltrano de Souza")
    salvar_modelo(config, caminho)


def salvar_modelo(config: LivroPontoConfig, caminho: str | Path) -> None:
    """Grava um LivroPontoConfig no formato .xlsx simplificado (abas Escola/
    Pessoas/Excecoes) — usado tanto por `criar-modelo` quanto pelo botão
    "Salvar" do app web."""
    wb = Workbook()
    escola = config.escola

    aba_escola = wb.active
    aba_escola.title = "Escola"
    aba_escola.append(["campo", "valor"])
    valores_escola = {
        "nome": escola.nome,
        "diretoria_ensino": escola.diretoria_ensino,
        "endereco": escola.endereco,
        "municipio": escola.municipio,
        "telefone1": escola.telefone1,
        "telefone2": escola.telefone2,
        "email": escola.email,
        "codigo_ua": escola.codigo_ua,
        "codigo_cie": escola.codigo_cie,
    }
    for campo, _rotulo in CAMPOS_ESCOLA:
        aba_escola.append([campo, valores_escola.get(campo, "")])
    aba_escola.append(["mes", config.mes])
    aba_escola.append(["ano", config.ano])
    aba_escola.append(["cidade_assinatura", config.cidade_assinatura])
    aba_escola.append(["diretor_nome", config.diretor_nome])
    aba_escola.append(["uf", config.uf])

    aba_pessoas = wb.create_sheet("Pessoas")
    aba_pessoas.append(COLUNAS_PESSOAS)
    for p in config.pessoas:
        aba_pessoas.append(_linha_pessoa(p))

    aba_excecoes = wb.create_sheet("Excecoes")
    aba_excecoes.append(COLUNAS_EXCECOES)
    for e in config.dias_excecao:
        aba_excecoes.append([e.mes, e.dia, e.tipo, e.descricao])

    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


def _valor(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return v


def ler_modelo(caminho: str | Path) -> LivroPontoConfig:
    wb = load_workbook(caminho, data_only=True)

    campos: dict[str, str] = {}
    aba_escola = wb["Escola"]
    for row in aba_escola.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        chave = str(row[0]).strip()
        valor = _valor(row[1]) if len(row) > 1 else ""
        campos[chave] = valor

    escola = Escola(
        nome=str(campos.get("nome", "")),
        diretoria_ensino=str(campos.get("diretoria_ensino", "")),
        endereco=str(campos.get("endereco", "")),
        municipio=str(campos.get("municipio", "")),
        telefone1=str(campos.get("telefone1", "")),
        telefone2=str(campos.get("telefone2", "")),
        email=str(campos.get("email", "")),
        codigo_ua=str(campos.get("codigo_ua", "")),
        codigo_cie=str(campos.get("codigo_cie", "")),
    )
    mes = int(campos["mes"])
    ano = int(campos["ano"])
    cidade_assinatura = str(campos.get("cidade_assinatura") or escola.municipio)
    diretor_nome = str(campos.get("diretor_nome", ""))
    uf = str(campos.get("uf") or "SP")

    def _ler_tabela(nome_aba: str) -> tuple[dict[str, int], list[tuple]]:
        aba = wb[nome_aba]
        linhas = list(aba.iter_rows(min_row=1, values_only=True))
        if not linhas:
            return {}, []
        header = [str(c).strip().lower() if c else "" for c in linhas[0]]
        idx = {n: i for i, n in enumerate(header)}
        return idx, linhas[1:]

    def campo(row, idx, nome, default=""):
        i = idx.get(nome)
        if i is None or i >= len(row):
            return default
        v = row[i]
        return default if v is None else v

    idx_pessoas, linhas_pessoas = _ler_tabela("Pessoas")
    pessoas: list[Pessoa] = []
    for row in linhas_pessoas:
        if not row or all(v in (None, "") for v in row):
            continue
        nome = str(campo(row, idx_pessoas, "nome", "")).strip()
        if not nome:
            continue
        tipo_raw = str(campo(row, idx_pessoas, "tipo", "ADMINISTRATIVO")).strip().upper()
        if tipo_raw.startswith("DOC"):
            tipo = TipoServidor.DOCENTE
        elif tipo_raw.startswith("GES"):
            tipo = TipoServidor.GESTAO
        else:
            tipo = TipoServidor.ADMINISTRATIVO
        ponto_raw = str(campo(row, idx_pessoas, "ponto", "S")).strip().upper()
        jornada_raw = campo(row, idx_pessoas, "jornada_semanal", "")
        pessoas.append(
            Pessoa(
                nome=nome,
                tipo=tipo,
                rg=str(campo(row, idx_pessoas, "rg", "")),
                cargo=str(campo(row, idx_pessoas, "cargo", "")),
                jornada_semanal=float(jornada_raw) if jornada_raw not in ("", None) else None,
                jornada_codigo=str(campo(row, idx_pessoas, "jornada_codigo", "")),
                ponto=(ponto_raw != "N"),
                entrada=str(campo(row, idx_pessoas, "entrada", "")),
                saida=str(campo(row, idx_pessoas, "saida", "")),
                intervalo_inicio=str(campo(row, idx_pessoas, "intervalo_inicio", "")),
                intervalo_fim=str(campo(row, idx_pessoas, "intervalo_fim", "")),
                ferias_inicio=str(campo(row, idx_pessoas, "ferias_inicio", "")),
                ferias_fim=str(campo(row, idx_pessoas, "ferias_fim", "")),
                disciplinas=str(campo(row, idx_pessoas, "disciplinas", "")),
                categoria=str(campo(row, idx_pessoas, "categoria", "")),
                situacao=str(campo(row, idx_pessoas, "situacao", "")),
                observacoes=str(campo(row, idx_pessoas, "observacoes", "")),
            )
        )

    dias_excecao: list[DiaNaoLetivo] = []
    if "Excecoes" in wb.sheetnames:
        idx_exc, linhas_exc = _ler_tabela("Excecoes")
        for row in linhas_exc:
            if not row or all(v in (None, "") for v in row):
                continue
            mes_e = campo(row, idx_exc, "mes", "")
            dia_e = campo(row, idx_exc, "dia", "")
            if mes_e in ("", None) or dia_e in ("", None):
                continue
            dias_excecao.append(
                DiaNaoLetivo(
                    mes=int(mes_e),
                    dia=int(dia_e),
                    tipo=str(campo(row, idx_exc, "tipo", "")).strip().upper(),
                    descricao=str(campo(row, idx_exc, "descricao", "")),
                )
            )

    return LivroPontoConfig(
        escola=escola,
        mes=mes,
        ano=ano,
        pessoas=pessoas,
        cidade_assinatura=cidade_assinatura,
        diretor_nome=diretor_nome,
        uf=uf,
        dias_excecao=dias_excecao,
    )
